from openpyxl import load_workbook
from pathlib import Path

RUTA = Path("datos") / "mayo2026.xlsx"

wb = load_workbook(RUTA, data_only=True)
hoja = wb["MAYO 2026"]

print("=" * 110)
print("INSPECCIÓN DETALLADA AE6, AE7, AE8")
print("=" * 110)


def volcar(fila_ini, fila_fin, etiqueta):
    print(f"\n--- {etiqueta} (filas {fila_ini}–{fila_fin}) ---\n")
    print(f"{'Fila':>5} | {'A':<8} | {'B':<6} | {'C':<6} | {'D':<10} | {'E':<32} | {'F':<5} | {'H':<5} | {'J':<5}")
    print("-" * 110)
    for fila in range(fila_ini, fila_fin + 1):
        a = hoja.cell(fila, 1).value
        b = hoja.cell(fila, 2).value
        c = hoja.cell(fila, 3).value
        d = hoja.cell(fila, 4).value
        e = hoja.cell(fila, 5).value
        f = hoja.cell(fila, 6).value
        h = hoja.cell(fila, 8).value
        j = hoja.cell(fila, 10).value
        # Solo imprimimos si hay algo relevante
        if a or d or e or f:
            print(f"{fila:>5} | {str(a or ''):<8} | {str(b or ''):<6} | {str(c or ''):<6} | {str(d or ''):<10} | {str(e or '')[:32]:<32} | {str(f or ''):<5} | {str(h or ''):<5} | {str(j or ''):<5}")


# Vamos a mirar unas filas antes para ver la cabecera
volcar(135, 180, "AE6 (con cabecera)")
volcar(193, 245, "AE7 (con cabecera)")
volcar(259, 300, "AE8 (con cabecera)")