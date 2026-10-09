"""
Volcado real de cambios aprobados a la hoja MAYO 2026 del Excel.
"""
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook

from app.db.modelos import Agente, Cuadrante

RUTA_EXCEL = Path("datos") / "mayo2026.xlsx"


def _leer_datos_para_cambio(db, solicitud):
    """
    Devuelve (fila_sol, col_sol, fila_comp, col_comp) o None si falta algo.
    """
    ag_sol = db.query(Agente).filter_by(id=solicitud.solicitante_id).first()
    ag_comp = db.query(Agente).filter_by(id=solicitud.companero_id).first()
    if not ag_sol or not ag_comp:
        return None

    reg_sol = db.query(Cuadrante).filter_by(
        agente_id=ag_sol.id, mes=solicitud.mes, anio=solicitud.anio,
        dia=solicitud.dia_solicitante).first()
    reg_comp = db.query(Cuadrante).filter_by(
        agente_id=ag_comp.id, mes=solicitud.mes, anio=solicitud.anio,
        dia=solicitud.dia_companero).first()

    if not reg_sol or not reg_comp:
        return None

    return (
        ag_sol.fila_excel, reg_sol.col_excel,
        ag_comp.fila_excel, reg_comp.col_excel,
    )


def volcar_a_excel(db, cambios):
    """
    cambios: lista de tuplas (CambioAplicado, Solicitud)
    Devuelve (ids_ok, errores)
    """
    if not cambios:
        return [], []

    if not RUTA_EXCEL.exists():
        return [], [f"No existe {RUTA_EXCEL}"]

    try:
        wb = load_workbook(RUTA_EXCEL)
    except Exception as e:
        return [], [f"No se pudo abrir {RUTA_EXCEL}: {e}"]

    if "MAYO 2026" not in wb.sheetnames:
        return [], ["La hoja 'MAYO 2026' no existe en el libro"]
    hoja = wb["MAYO 2026"]

    ids_ok = []
    errores = []

    for cambio, solicitud in cambios:
        try:
            datos = _leer_datos_para_cambio(db, solicitud)
            if not datos:
                errores.append(f"Solicitud {solicitud.id}: faltan datos de agente/cuadrante")
                continue

            fila_sol, col_sol, fila_comp, col_comp = datos

            # Intercambio: A recibe el servicio de B, y B recibe el de A
            hoja.cell(row=fila_sol, column=col_sol).value = solicitud.servicio_companero
            hoja.cell(row=fila_comp, column=col_comp).value = solicitud.servicio_solicitante

            ids_ok.append(cambio.id)
        except Exception as e:
            errores.append(f"Cambio {cambio.id} (solicitud {solicitud.id}): {e}")

    # Guardar solo si algo se ha tocado
    if ids_ok:
        try:
            wb.save(RUTA_EXCEL)
        except Exception as e:
            return [], [f"No se pudo guardar {RUTA_EXCEL}: {e}"]

    return ids_ok, errores